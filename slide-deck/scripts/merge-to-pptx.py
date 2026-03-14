#!/usr/bin/env python3
"""
Merge slide images into a single PPTX presentation.

Usage:
  python3 merge-to-pptx.py <slide-deck-dir> [--output filename.pptx]

Finds images matching NN-slide-*.{png,jpg,jpeg} and creates a 16:9 PPTX
with each image as a full-cover slide. Optionally adds speaker notes from
prompts/ directory.

Requires: python-pptx, Pillow
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu


# ── Constants ─────────────────────────────────────────

SLIDE_WIDTH_EMU = Emu(12192000)   # 13.333 inches (16:9)
SLIDE_HEIGHT_EMU = Emu(6858000)   # 7.5 inches

SLIDE_PATTERN = re.compile(r'^(\d+)-slide-.*\.(png|jpg|jpeg)$', re.IGNORECASE)

SCRIPT_DIR = Path(__file__).resolve().parent


# ── Slide discovery ───────────────────────────────────

def find_slide_images(deck_dir: Path) -> list[dict]:
    """Find and sort slide images by their numeric prefix."""
    if not deck_dir.is_dir():
        print(f"Directory not found: {deck_dir}", file=sys.stderr)
        sys.exit(1)

    slides = []
    prompts_dir = deck_dir / "prompts"

    for f in deck_dir.iterdir():
        m = SLIDE_PATTERN.match(f.name)
        if not m:
            continue

        stem = f.stem  # e.g. "01-slide-title"
        prompt_path = prompts_dir / f"{stem}.md"

        slides.append({
            "filename": f.name,
            "path": f,
            "index": int(m.group(1)),
            "prompt_path": prompt_path if prompt_path.exists() else None,
        })

    slides.sort(key=lambda s: s["index"])

    if not slides:
        print(f"No slide images found in: {deck_dir}", file=sys.stderr)
        print("Expected format: 01-slide-*.png, 02-slide-*.png, etc.", file=sys.stderr)
        sys.exit(1)

    return slides


def find_base_prompt() -> str | None:
    """Read references/base-prompt.md relative to this script."""
    base_prompt_path = SCRIPT_DIR / ".." / "references" / "base-prompt.md"
    if base_prompt_path.exists():
        return base_prompt_path.read_text(encoding="utf-8")
    return None


# ── PPTX creation ─────────────────────────────────────

def create_pptx(slides: list[dict], output_path: Path) -> None:
    """Build PPTX with full-cover slide images and optional speaker notes."""
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH_EMU
    prs.slide_height = SLIDE_HEIGHT_EMU
    prs.core_properties.author = "slide-deck"
    prs.core_properties.subject = "Generated Slide Deck"

    blank_layout = prs.slide_layouts[6]  # Blank
    base_prompt = find_base_prompt()
    notes_count = 0

    for slide_info in slides:
        slide = prs.slides.add_slide(blank_layout)

        # Add image as full-cover background
        slide.shapes.add_picture(
            str(slide_info["path"]),
            Emu(0), Emu(0),
            SLIDE_WIDTH_EMU, SLIDE_HEIGHT_EMU,
        )

        # Add speaker notes from prompt file
        if slide_info["prompt_path"]:
            slide_prompt = slide_info["prompt_path"].read_text(encoding="utf-8")
            if base_prompt:
                full_notes = f"{base_prompt}\n\n---\n\n{slide_prompt}"
            else:
                full_notes = slide_prompt

            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = full_notes
            notes_count += 1

        status = " (with notes)" if slide_info["prompt_path"] else ""
        print(f"Added: {slide_info['filename']}{status}")

    prs.save(str(output_path))
    print(f"\nCreated: {output_path}")
    print(f"Total slides: {len(slides)}")
    if notes_count > 0:
        base_note = " (includes base prompt)" if base_prompt else ""
        print(f"Slides with notes: {notes_count}{base_note}")


# ── Main ──────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(description="Merge slide images into PPTX")
    parser.add_argument("dir", help="Slide deck directory containing NN-slide-*.{png,jpg,jpeg}")
    parser.add_argument("--output", "-o", help="Output PPTX filename (default: <dir>/<dirName>.pptx)")
    args = parser.parse_args()

    deck_dir = Path(args.dir).resolve()
    slides = find_slide_images(deck_dir)

    # Default output: {dir}/{dirName}.pptx
    if args.output:
        output_path = Path(args.output)
    else:
        dir_name = deck_dir.name
        if dir_name == "slide-deck":
            dir_name = deck_dir.parent.name
        output_path = deck_dir / f"{dir_name}.pptx"

    print(f"Found {len(slides)} slides in: {deck_dir}\n")

    create_pptx(slides, output_path)


if __name__ == "__main__":
    main()
