"""PPTX construction: transparent textboxes + precise font sizing + multi-color runs."""
from __future__ import annotations

from math import ceil
from typing import Optional

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

from .common import (
    SlideArtifacts, TextRegion, TextStyle, ColoredSegment,
    SLIDE_WIDTH_INCHES, SLIDE_HEIGHT_INCHES, PPI, ALIGN_MAP,
    get_default_font_name, hex_to_rgb, find_cjk_font_path,
)

EXPAND_RATIO = 0.01  # 1% bbox expansion to prevent text clipping
MAX_FONT_SIZE = 200
MIN_FONT_SIZE = 6

# Font cache for text measurement
_font_cache: dict[tuple[str, int], any] = {}
_cjk_font_path: Optional[str] = None


def _get_font(size_pt: int) -> Optional[ImageFont.FreeTypeFont]:
    """Get a Pillow ImageFont for text measurement."""
    global _cjk_font_path
    if _cjk_font_path is None:
        _cjk_font_path = find_cjk_font_path() or ""
    if not _cjk_font_path:
        return None
    key = (_cjk_font_path, size_pt)
    if key not in _font_cache:
        try:
            _font_cache[key] = ImageFont.truetype(_cjk_font_path, size_pt)
        except Exception:
            return None
    return _font_cache[key]


def _measure_text_width(text: str, font_size_pt: int) -> Optional[float]:
    """Measure text width in points using Pillow."""
    font = _get_font(font_size_pt)
    if font is None:
        return None
    bbox = font.getbbox(text)
    width_px = bbox[2] - bbox[0]
    width_pt = width_px * 72.0 / PPI
    return width_pt


def estimate_font_size(text: str, bbox_w_px: float, bbox_h_px: float) -> float:
    """Estimate the largest font size that fits in the bbox.

    Banana-slides approach: descending scan with actual measurement.
    """
    w_pt = bbox_w_px * 72 / PPI
    h_pt = bbox_h_px * 72 / PPI
    lines = text.split("\n")

    best = MIN_FONT_SIZE
    for fs in range(MAX_FONT_SIZE, MIN_FONT_SIZE - 1, -1):
        total_height = 0.0
        for line in lines:
            measured = _measure_text_width(line, fs) if line else None
            if measured is not None:
                lines_needed = max(1, ceil(measured / w_pt)) if w_pt > 0 else 1
            else:
                # Heuristic: CJK chars ~1.0em, latin ~0.6em
                char_width = sum(
                    1.0 if ord(c) > 0x2E80 else 0.6 for c in line
                ) * fs if line else 0
                lines_needed = max(1, ceil(char_width / w_pt)) if w_pt > 0 else 1
            total_height += lines_needed * fs
        if total_height <= h_pt:
            best = fs
            break

    return float(max(MIN_FONT_SIZE, min(MAX_FONT_SIZE, best)))


def _add_text_element(slide, region: TextRegion, scale_x: float, scale_y: float, prs) -> None:
    """Add one text region to a PPTX slide."""
    left_px = region.left * scale_x
    top_px = region.top * scale_y
    w_px = region.width * scale_x
    h_px = region.height * scale_y

    # Expand bbox by EXPAND_RATIO to prevent clipping
    expand_w = w_px * EXPAND_RATIO
    expand_h = h_px * EXPAND_RATIO
    left_px = max(0, left_px - expand_w / 2)
    top_px = max(0, top_px - expand_h / 2)
    w_px += expand_w
    h_px += expand_h

    left = Inches(left_px / PPI)
    top = Inches(top_px / PPI)
    width = Inches(w_px / PPI)
    height = Inches(h_px / PPI)

    txBox = slide.shapes.add_textbox(left, top, width, height)

    # Transparent fill (default) + no border
    txBox.line.fill.background()

    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)

    font_size = estimate_font_size(region.text, w_px, h_px)

    style = region.style or TextStyle()
    alignment = ALIGN_MAP.get(style.align, ALIGN_MAP["left"])

    if style.colored_segments and len(style.colored_segments) > 0:
        # Multi-color path
        p = tf.paragraphs[0]
        p.clear()
        p.alignment = alignment
        for seg in style.colored_segments:
            run = p.add_run()
            run.text = seg.text
            run.font.size = Pt(font_size)
            run.font.name = get_default_font_name(seg.text)
            run.font.color.rgb = RGBColor(*hex_to_rgb(seg.color))
            run.font.bold = style.bold
            if style.italic:
                run.font.italic = True
    else:
        # Single-color path
        text = region.text
        lines = text.split("\n")
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(font_size)
            p.font.name = get_default_font_name(line or text)
            p.font.color.rgb = RGBColor(*hex_to_rgb(style.color))
            p.font.bold = style.bold
            if style.italic:
                p.font.italic = True
            p.alignment = alignment


def build_presentation(slides_data: list[SlideArtifacts], output_path: str) -> None:
    """Build a complete PPTX from processed slide data."""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_INCHES)
    prs.slide_height = Inches(SLIDE_HEIGHT_INCHES)

    for sa in slides_data:
        slide_layout = prs.slide_layouts[6]  # blank
        slide = prs.slides.add_slide(slide_layout)

        img_w, img_h = sa.image_size
        slide_w_px = SLIDE_WIDTH_INCHES * PPI
        slide_h_px = SLIDE_HEIGHT_INCHES * PPI
        scale_x = slide_w_px / img_w
        scale_y = slide_h_px / img_h

        # Add clean background as full-slide picture
        slide.shapes.add_picture(
            sa.clean_bg_path, Emu(0), Emu(0),
            prs.slide_width, prs.slide_height,
        )

        for region in sa.regions:
            _add_text_element(slide, region, scale_x, scale_y, prs)

    prs.save(output_path)
    print(f"✓ 可编辑 PPTX 已保存: {output_path}")
